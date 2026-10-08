"""استخراج النص من PDF و Word."""
from __future__ import annotations

import io
import logging
import re
import shutil
import subprocess
import uuid
from pathlib import Path

from .errors import CorruptedFileError
from .bidi import Glyph, glyphs_to_line_records, glyphs_to_lines
from .layout import Block

try:
    import pymupdf as fitz
except ImportError:  # إصدارات قديمة
    import fitz  # type: ignore

log = logging.getLogger("mithaq.docproc")


# ------------------------------- PDF ---------------------------------------
def _arabic_dominant(text: str) -> bool:
    return len(re.findall(r"[\u0600-\u06FF]", text)) > len(re.findall(r"[A-Za-z]", text))


_FONT_GLYPH_FIXES = {
    "GESS": {"‡": "لأ", "—": "لأ", "Ë": "لأ", "Œ": "لإ", "‰": "لإ", "é": "لإ", "Õ": "لآ",
             "Ï": "اً", "(cid:15)": "اً", "\x0f": "اً", "\ufffd": "اً"},
}


def _glyph_fix_table(font: str) -> dict | None:
    key = (font or "").replace(" ", "").replace("-", "").upper()
    for name, table in _FONT_GLYPH_FIXES.items():
        if name in key:
            return table
    return None


def _page_glyphs(page) -> list[list[Glyph]]:
    out, seq = [], 0
    for b in page.get_text("rawdict", sort=False)["blocks"]:
        if b.get("type") != 0:
            continue
        gs = []
        for line in b["lines"]:
            for s in line["spans"]:
                fix = _glyph_fix_table(s.get("font", ""))
                for ch in s["chars"]:
                    x0, y0, x1, y1 = ch["bbox"]
                    c = ch["c"]
                    if fix and c in fix:
                        c = fix[c]
                    gs.append(Glyph(c, x0, y0, x1, y1, seq))
                    seq += 1
        if gs:
            out.append(gs)
    return out


def merge_split_blocks(groups: list[list[Glyph]]) -> list[list[Glyph]]:
    
    n = len(groups)
    solid = [[g for g in gs if not g.text.isspace()] for gs in groups]
    yr = [(min(g.y0 for g in s), max(g.y1 for g in s)) if s else None for s in solid]
    parent = list(range(n))

    def find(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    def touching(A, B) -> bool:
        lo, hi = max(yr[A][0], yr[B][0]), min(yr[A][1], yr[B][1])
        if hi <= lo:
            return False
        ga = sorted((g for g in solid[A] if lo - 1 <= g.cy <= hi + 1), key=lambda g: g.x0)
        gb = sorted((g for g in solid[B] if lo - 1 <= g.cy <= hi + 1), key=lambda g: g.x0)
        if not ga or not gb:
            return False
        import bisect
        xs = [g.x0 for g in gb]
        wmax = max(g.x1 - g.x0 for g in gb)
        for a in ga:
            h = a.h
            i = bisect.bisect_left(xs, a.x0 - wmax - h)
            for b in gb[i:]:
                if b.x0 > a.x1 + 0.5 * h:
                    break
                if abs(a.cy - b.cy) < 0.3 * max(a.h, b.h) and (max(a.x0, b.x0) - min(a.x1, b.x1)) < 0.5 * max(a.h, b.h):
                    return True
        return False

    for i in range(n):
        if not yr[i]:
            continue
        for j in range(i + 1, n):
            if yr[j] and find(i) != find(j) and touching(i, j):
                parent[find(j)] = find(i)
    merged: dict[int, list[Glyph]] = {}
    order: list[int] = []
    for i in range(n):
        r = find(i)
        if r not in merged:
            merged[r] = []
            order.append(r)
        merged[r].extend(groups[i])
    return [merged[r] for r in order]


def _inside(g: Glyph, r) -> bool:
    return r[0] <= g.cx <= r[2] and r[1] <= g.cy <= r[3]


def _bbox_of(gs: list[Glyph]) -> tuple:
    return (min(g.x0 for g in gs), min(g.y0 for g in gs), max(g.x1 for g in gs), max(g.y1 for g in gs))


def _clean_rows(rows: list[list[str]]) -> list[list[str]] | None:
    from collections import Counter
    rows = [r for r in rows if any(r)]
    if not rows:
        return None
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    keep = [j for j in range(width) if any(r[j] for r in rows)]          # نحذف الأعمدة الفاضية كلياً
    rows = [[r[j] for j in keep] for r in rows]
    compact = [[c for c in r if c] for r in rows]
    cnt = Counter(len(c) for c in compact if len(c) >= 2)
    if cnt and cnt.most_common(1)[0][1] >= 0.6 * sum(1 for c in compact if len(c) >= 2):
        rows = compact
    if len(rows) < 2 or len(keep) < 2:
        return None                                                     # غالباً مربع نص وليس جدول
    n_cols = max(len(r) for r in rows)
    fill = sum(1 for r in rows for c in r if c) / (len(rows) * n_cols)
    if n_cols > 6 and fill < 0.3:
        return None                                                     # شبكة تنسيق للصفحة وليست جدول بيانات
    return rows


def _split_paragraphs(gs: list[Glyph]) -> list[Block]:

    recs = glyphs_to_line_records(gs)
    if not recs:
        return []
    bx0, by0, bx1, by1 = _bbox_of(gs)
    width = max(bx1 - bx0, 1.0)
    rtl = _arabic_dominant(" ".join(r[0] for r in recs))
    paras, cur = [], []
    for k, rec in enumerate(recs):
        cur.append(rec)
        t, x0, x1 = rec
        short = (x0 - bx0 > 0.15 * width) if rtl else (bx1 - x1 > 0.15 * width)
        if short and k < len(recs) - 1:
            paras.append(cur)
            cur = []
    if cur:
        paras.append(cur)
    out, n, i = [], len(recs), 0
    h = (by1 - by0) / n
    for p in paras:
        out.append(Block("\n".join(r[0] for r in p),
                         (min(r[1] for r in p), by0 + i * h, max(r[2] for r in p), by0 + (i + len(p)) * h)))
        i += len(p)
    return out


def page_glyph_boxes(page) -> list[tuple]:
    return [(g.x0, g.y0, g.x1, g.y1) for gs in _page_glyphs(page) for g in gs if not g.text.isspace()]


def native_blocks(page, extra_glyphs: list[Glyph] | None = None) -> list[Block]:
  
    groups = merge_split_blocks(_page_glyphs(page))
    if extra_glyphs:                                   # نص مقروء بـ OCR من مناطق مرسومة بدون طبقة نص
        groups += [[g] for g in extra_glyphs]
    all_glyphs = [g for gs in groups for g in gs]

    table_rects, table_blocks = [], []
    owned_ids: set[int] = set()                  # حروف دخلت خلية جدول (الباقي يبقى في النص العادي)
    try:
        tables = page.find_tables().tables
    except Exception as e:                                   # find_tables غير متوفر أو فشل
        log.debug("find_tables skipped: %s", e)
        tables = []
    for t in tables:
        try:
            
            boxes = [cb for row in t.rows for cb in row.cells if cb is not None]
            area = lambda r: (r[2] - r[0]) * (r[3] - r[1])
            owner: dict[int, tuple] = {}
            tb = tuple(t.bbox)
            for g in all_glyphs:
                if not _inside(g, tb):
                    continue
                hit = [cb for cb in boxes if _inside(g, cb)]
                if hit:
                    owner[id(g)] = tuple(min(hit, key=area))
            rows = []
            for row in t.rows:
                cells = []
                for cb in row.cells:
                    if cb is None:                                # جزء من خلية مدموجة
                        cells.append("")
                        continue
                    gs = [g for g in all_glyphs if owner.get(id(g)) == tuple(cb)]
                    cells.append(" ".join(glyphs_to_lines([Glyph(g.text, g.x0, g.y0, g.x1, g.y1, g.seq) for g in gs])))
                rows.append(cells)
            rows = _clean_rows(rows)
            if rows is None:
                continue
            if _arabic_dominant(" ".join(" ".join(r) for r in rows)):
                rows = [r[::-1] for r in rows]                    # الجدول العربي: العمود الأول على اليمين
            rect = tuple(t.bbox)
            table_rects.append(rect)
            owned_ids.update(owner.keys())
            table_blocks.append(Block("\n".join(" | ".join(r) for r in rows), rect, "table", rows=rows))
        except Exception as e:
            log.debug("table skipped: %s", e)

    blocks: list[Block] = []
    for gs in groups:
        gs = [g for g in gs if id(g) not in owned_ids]            # ما نحذف نص داخل حدود الجدول وبرّا خلاياه
        if not gs or not "".join(g.text for g in gs).strip():
            continue
        for para in _split_paragraphs(gs):
            blocks.append(para)
    blocks.extend(table_blocks)
    rtl = _arabic_dominant("".join(g.text for g in all_glyphs))
    blocks.sort(key=lambda b: (round(b.bbox[1]), -b.bbox[2] if rtl else b.bbox[0]))   # من الأعلى، ثم حسب اتجاه الصفحة
    return blocks


def uncovered_ink_regions(page, glyph_boxes: list[tuple], dpi: int = 100,
                          ink_threshold: int = 150, min_h_pt: float = 4.0, min_w_pt: float = 6.0) -> list[tuple]:
    
    import numpy as np
    from PIL import Image
    img = render_page(page, dpi, 10 ** 6).convert("L")
    a = np.asarray(img)
    sx, sy = img.width / page.rect.width, img.height / page.rect.height
    ink = a < ink_threshold
    dark_orig = ink.copy()                                 # قبل أي حذف: لكشف الخلفيات الغامقة
    pad = max(1, int(round(1.5 * sx)))
    try:                                                   # الصور (شعارات، أختام) مو نص مرسوم
        image_boxes = [tuple(i["bbox"]) for i in page.get_image_info()]
    except Exception:
        image_boxes = []
    for (x0, y0, x1, y1) in image_boxes:
        ink[max(0, int(y0 * sy)):int(y1 * sy) + 1, max(0, int(x0 * sx)):int(x1 * sx) + 1] = False
    for (x0, y0, x1, y1) in glyph_boxes:                 # نغطي أماكن الحروف المعروفة
        ink[max(0, int(y0 * sy) - pad):int(y1 * sy) + pad + 1, max(0, int(x0 * sx) - pad):int(x1 * sx) + pad + 1] = False
    
    def _erase_runs(m, limit):                             # m: مصفوفة؛ نمسح الامتدادات الطويلة على المحور 0
        for j in range(m.shape[1]):
            col = m[:, j]
            if col.sum() < limit:
                continue
            d = np.diff(np.concatenate(([0], col.view(np.int8), [0])))
            starts, ends = np.where(d == 1)[0], np.where(d == -1)[0]
            for s, t in zip(starts, ends):
                if t - s >= limit:
                    col[s:t] = False
    _erase_runs(ink, int(25 * sy))
    _erase_runs(ink.T, int(60 * sx))
    rows = ink.sum(axis=1)
    min_h, min_w = int(min_h_pt * sy), int(min_w_pt * sx)
    regions = []
    y = 0
    H = ink.shape[0]
    while y < H:                                          # أشرطة أفقية (أسطر)
        if rows[y] == 0:
            y += 1
            continue
        y0 = y
        while y < H and rows[y] > 0:
            y += 1
        band = ink[y0:y]
        if y - y0 < min_h or y - y0 > 40 * sy:            # أطول من سطرين-ثلاثة = رسم/صورة، مو سطر نص
            continue
        cols = band.sum(axis=0) > 0
        x, W = 0, band.shape[1]
        gap = int(3 * (y - y0))                             # فجوة أكبر من 3 ارتفاعات سطر = عمود آخر
        while x < W:
            if not cols[x]:
                x += 1
                continue
            x0 = x
            last = x
            while x < W and (cols[x] or x - last <= gap):
                if cols[x]:
                    last = x
                x += 1
            x1 = last + 1
            seg = band[:, x0:x1]
            fill = seg.mean()
            bg_dark = dark_orig[y0:y, x0:x1].mean()          # خلفية غامقة بنص أبيض (رؤوس جداول) = مو نص مفقود
            if x1 - x0 >= min_w and 0.03 < fill < 0.6 and bg_dark < 0.45:
                regions.append((x0 / sx, y0 / sy, x1 / sx, y / sy))
    return regions


def render_page(page, dpi: int, max_side_px: int):
    from PIL import Image
    w, h = page.rect.width, page.rect.height
    zoom = min(dpi / 72.0, max_side_px / max(w, h))
    pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
    return Image.frombytes("RGB", (pix.width, pix.height), pix.samples)


# ------------------------------- DOCX --------------------------------------
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_A_BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
_R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"


def docx_direct_blocks(path) -> tuple[list[Block], dict]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    try:
        doc = Document(str(path))
    except Exception as e:
        raise CorruptedFileError(f"تعذر فتح ملف Word: {e}") from e

    blocks: list[Block] = []
    numbered = 0
    for child in doc.element.body.iterchildren():
        if child.tag == _W + "p":
            p = Paragraph(child, doc)
            if p._p.pPr is not None and p._p.pPr.numPr is not None:
                numbered += 1
            if p.text.strip():
                blocks.append(Block(p.text, (0, 0, 0, 0)))
        elif child.tag == _W + "tbl":
            t = Table(child, doc)
            rows = []
            for row in t.rows:
                seen, cells = set(), []
                for c in row.cells:
                    if id(c._tc) in seen:                    # الخلايا المدموجة تتكرر
                        continue
                    seen.add(id(c._tc))
                    cells.append(" ".join(c.text.split()))
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                grid = [r.split(" | ") for r in rows]
                blocks.append(Block("\n".join(rows), (0, 0, 0, 0), "table", rows=grid))
    return blocks, {"numbered_paragraphs": numbered}


def docx_images(path) -> list[bytes]:
    from docx import Document
    doc = Document(str(path))
    out = []
    for blip in doc.element.body.iter(_A_BLIP):
        rid = blip.get(_R_EMBED)
        part = doc.part.related_parts.get(rid) if rid else None
        if part is not None and part.content_type.startswith("image/"):
            out.append(part.blob)
    return out


def _find_soffice() -> str | None:
    import os
    for c in (os.environ.get("LIBREOFFICE_PATH"), shutil.which("soffice"), shutil.which("libreoffice"),
              r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if c and os.path.isfile(c):
            return c
    return None


def convert_docx_to_pdf(src, outdir, timeout: int = 120) -> Path | None:
    exe = _find_soffice()
    if not exe:
        return None
    outdir = Path(outdir)
    
    profile = (outdir / f"lo_profile_{uuid.uuid4().hex[:8]}").resolve()
    try:
        subprocess.run([exe, f"-env:UserInstallation={profile.as_uri()}", "--headless",
                        "--convert-to", "pdf", "--outdir", str(outdir), str(src)],
                       capture_output=True, timeout=timeout, check=True)
    except Exception as e:
        log.warning("LibreOffice conversion failed: %s", e)
        return None
    pdf = outdir / (Path(src).stem + ".pdf")
    return pdf if pdf.exists() and pdf.stat().st_size > 0 else None


def _skeleton(text: str) -> set:
    from collections import Counter  # noqa: F401
    from .cleaning import normalize_for_search
    t = normalize_for_search(text).lower()
    t = re.sub("[ال]", "", t)
    return set(re.findall(r"[^\W\d_]{2,}", t))


def map_blocks_to_pages(blocks: list[Block], pdf_path) -> list[int]:
    doc = fitz.open(str(pdf_path))
    bags = []
    for p in doc:
        # نقرأ الصفحة بطريقتنا (ترتيب هندسي) لأن get_text() البسيط يطلّع العربي مقلوب فما يتطابق شي
        try:
            text = " ".join(b.text for b in native_blocks(p))
        except Exception:
            text = p.get_text()
        bags.append(_skeleton(text))
    doc.close()
    prev, result = 0, []
    for b in blocks:
        ws = _skeleton(b.text)
        best, best_score = prev, 0.3                     # أقل من 0.3 = ما فيه دليل، نبقى على الصفحة السابقة
        if ws:
            for pi in range(prev, min(prev + 3, len(bags))):
                score = len(ws & bags[pi]) / len(ws)
                if score > best_score:
                    best, best_score = pi, score
        prev = best
        result.append(best + 1)
    return result
