"""تشغيل كامل: Word (بدون PyMuPDF/LibreOffice) + ملفات PDF الحقيقية لو موجودة في testset/."""
import dataclasses
import json
import tempfile
from pathlib import Path

from mithaq_docproc import CONFIG, DocProcError, process_document, write_outputs

ROOT = Path(__file__).resolve().parents[2]
CFG = dataclasses.replace(CONFIG, ocr_engine="none", docx_page_numbers=False)


def _make_docx(path):
    from docx import Document
    d = Document()
    d.add_paragraph("الشروط والأحكام")
    d.add_paragraph("1 اسم الصندوق ونوعه")
    d.add_paragraph("صندوق أصول وبخيت للتمويل المباشر هو صندوق استثمار تمويل مباشر مغلق المدة ومطروح طرحاً خاصاً.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "رسوم الاشتراك", "%1.0 من قيمة الاشتراك"
    t.cell(1, 0).text, t.cell(1, 1).text = "رسوم أمين الحفظ", "%0.05 سنويا"
    d.save(path)


def test_docx_full_output():
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "sample.docx"
        _make_docx(src)
        res = process_document(src, cfg=CFG)
        types = [it["type"] for it in res.items]
        assert "table" in types and "heading" in types
        table = next(it for it in res.items if it["type"] == "table")
        assert table["table"]["rows"][0][0] == "رسوم الاشتراك"
        paths = write_outputs(res, Path(tmp) / "out", src)
        data = json.loads(paths["json"].read_text(encoding="utf-8"))
        assert data["schema_version"] == "1.0" and data["document"]["sha256"]
        assert {"json", "txt", "md"} <= set(paths)
        assert "\r" not in paths["txt"].read_text(encoding="utf-8")


def test_invalid_files_rejected():
    with tempfile.TemporaryDirectory() as tmp:
        bad = Path(tmp) / "x.pdf"
        bad.write_bytes(b"not a pdf")
        for p, code in ((bad, "CORRUPTED_FILE"), (Path(tmp) / "x.txt", None)):
            if code is None:
                p.write_text("hello")
            try:
                process_document(p, cfg=CFG)
                raise AssertionError("expected error")
            except DocProcError as e:
                assert code is None or e.code == code


def test_real_pdfs_if_available():
    """ضعوا test.pdf (أصول وبخيت) في testset/ لتشغيل هذا الاختبار."""
    pdf = ROOT / "testset" / "test.pdf"
    if not pdf.exists():
        return
    res = process_document(pdf, cfg=CFG)
    page23 = next(p for p in res.pages if p.page == 23)
    assert "رسوم الاشتراك" in page23.text
    assert "REVERSED_ARABIC" not in res.warnings
