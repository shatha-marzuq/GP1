import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from segmentation import segment_clauses

# so we can import the docproc code from the other folder
this_folder = os.path.dirname(os.path.abspath(__file__))
sys.path.append(os.path.join(this_folder, "..", "docproc"))

# how to run:
#   python run_segmentation.py terms.pdf          -> OCR + segmentation
#   python run_segmentation.py terms.json         -> segmentation only (JSON that docproc already made)
#   python run_segmentation.py a.pdf b.pdf c.pdf  -> many files at once, each in its own folder
# the results are saved in backend/segmentation_output/<file name>/


# step 1: get the items from the file
def get_items(file_path):
    if file_path.lower().endswith(".json"):
        # docproc already ran on this file, just read its JSON
        with open(file_path, encoding="utf-8") as f:
            data = json.load(f)
        return data["items"], data.get("document", {})

    # PDF or Word: run Shatha's docproc first
    from mithaq_docproc import DocProcError, build_output, process_document

    print("running OCR / document processing ...")
    try:
        result = process_document(file_path)
    except DocProcError as e:
        print("error:", e.code, e.message_ar)
        return None, None

    data = build_output(result, file_path)
    return data["items"], data["document"]


# step 2: save the clauses as JSON (this is what the system will use)
def save_json(clauses, document_info, path):
    output = {
        "document": document_info,
        "clause_count": len(clauses),
        "clauses": clauses,
    }
    with open(path, "w", encoding="utf-8") as f:
        json.dump(output, f, ensure_ascii=False, indent=2)


# step 3: save the clauses as Excel (easy to read, one clause per row)
def save_excel(clauses, path):
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Clauses"
    sheet.sheet_view.rightToLeft = True  # Arabic

    headers = ["Clause ID", "Type", "Clause no.", "Title", "Page from", "Page to",
               "Part", "Words", "Has table", "From OCR", "Text"]
    sheet.append(headers)

    # make the header bold with a color
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")

    for c in clauses:
        part = str(c["part"]) + "/" + str(c["parts"])
        sheet.append([
            c["clause_id"],
            c["section_type"],
            c["clause_number"],
            c["section_title"],
            c["page_start"],
            c["page_end"],
            part,
            c["word_count"],
            "yes" if c["has_table"] else "",
            "yes" if c["from_ocr"] else "",
            c["content"],
        ])

    # column widths
    widths = [10, 8, 10, 35, 10, 10, 7, 8, 10, 10, 90]
    letters = "ABCDEFGHIJK"
    for i in range(len(widths)):
        sheet.column_dimensions[letters[i]].width = widths[i]

    # wrap long text so the cell shows all of it
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    sheet.freeze_panes = "A2"  # keep the header visible when scrolling
    wb.save(path)


# run one file and save its JSON + Excel in its own folder
def run_one_file(file_path, out_folder):
    print()
    print("====", file_path)
    items, document_info = get_items(file_path)
    if items is None:
        return
    print("items from docproc:", len(items))

    clauses = segment_clauses(items)
    print("clauses after segmentation:", len(clauses))

    # every file gets its own folder, example: terms.pdf -> segmentation_output/terms/
    name = os.path.splitext(os.path.basename(file_path))[0]
    out_folder = os.path.join(out_folder, name)
    os.makedirs(out_folder, exist_ok=True)
    json_path = os.path.join(out_folder, name + "_clauses.json")
    excel_path = os.path.join(out_folder, name + "_clauses.xlsx")
    save_json(clauses, document_info, json_path)
    save_excel(clauses, excel_path)
    print("saved:", json_path)
    print("saved:", excel_path)


def main():
    if len(sys.argv) < 2:
        print("usage: python run_segmentation.py <file1> <file2> ...   (pdf, docx or docproc json)")
        sys.exit(1)

    out_folder = os.path.join(this_folder, "segmentation_output")
    os.makedirs(out_folder, exist_ok=True)

    for file_path in sys.argv[1:]:
        if not os.path.exists(file_path):
            print("file not found:", file_path)
            continue
        run_one_file(file_path, out_folder)


if __name__ == "__main__":
    main()
